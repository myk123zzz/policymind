import io
from pathlib import Path

from docx import Document as DocxDocument

from policymind.documents.models import DocumentBlock, ParsedDocument


class DocxParser:
    supported_mime_types: frozenset[str] = frozenset(
        {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
    )

    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        doc = DocxDocument(io.BytesIO(content))
        title = Path(filename).stem
        blocks: list[DocumentBlock] = []

        current_heading: tuple[str, ...] = ()
        paragraphs_text: list[str] = []

        for para in doc.paragraphs:
            style = para.style.name if para.style else ""
            text = para.text

            if not text.strip():
                if paragraphs_text:
                    blocks.append(
                        DocumentBlock(
                            text="\n".join(paragraphs_text),
                            page_number=1,
                            block_type="text",
                            heading_path=current_heading,
                        )
                    )
                    paragraphs_text.clear()
                continue

            if style.startswith("Heading"):
                if paragraphs_text:
                    blocks.append(
                        DocumentBlock(
                            text="\n".join(paragraphs_text),
                            page_number=1,
                            block_type="text",
                            heading_path=current_heading,
                        )
                    )
                    paragraphs_text.clear()
                current_heading = (text,)
                if not title or title == Path(filename).stem:
                    title = text
            else:
                paragraphs_text.append(text)

        if paragraphs_text:
            blocks.append(
                DocumentBlock(
                    text="\n".join(paragraphs_text),
                    page_number=1,
                    block_type="text",
                    heading_path=current_heading,
                )
            )

        return ParsedDocument(title=title, page_count=1, blocks=blocks)
