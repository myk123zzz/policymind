import io
from pathlib import Path

import openpyxl  # type: ignore[import-untyped]

from policymind.documents.models import DocumentBlock, ParsedDocument


class XlsxParser:
    supported_mime_types: frozenset[str] = frozenset(
        {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
    )

    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        title = Path(filename).stem
        blocks: list[DocumentBlock] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            # Build markdown-like table
            lines = [f"## Sheet: {sheet_name}", ""]
            header = rows[0]
            lines.append("| " + " | ".join(str(c or "") for c in header) + " |")
            lines.append("|" + "|".join("------" for _ in header) + "|")
            for row in rows[1:]:
                lines.append("| " + " | ".join(str(c or "") for c in row) + " |")

            blocks.append(
                DocumentBlock(
                    text="\n".join(lines),
                    page_number=1,
                    block_type="table",
                )
            )

        wb.close()
        return ParsedDocument(title=title, page_count=1, blocks=blocks)
