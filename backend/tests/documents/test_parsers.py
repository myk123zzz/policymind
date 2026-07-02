import io


def test_markdown_parser_extracts_headings() -> None:
    """Markdown 解析器提取标题层级。"""
    from policymind.documents.parsers.markdown import MarkdownParser

    content = b"# Title\n\n## Section 1\n\nParagraph text here.\n\n## Section 2\n\nMore text."
    parser = MarkdownParser()
    result = parser.parse(content, filename="test.md")

    assert result.title == "Title"
    assert result.page_count == 1
    assert len(result.blocks) == 2


def test_markdown_parser_title_fallback() -> None:
    """无标题时用文件名作为回退。"""
    from policymind.documents.parsers.markdown import MarkdownParser

    content = b"Just some text\n\nwithout any heading."
    parser = MarkdownParser()
    result = parser.parse(content, filename="notes.md")

    assert result.title is not None


def test_pdf_parser_extracts_text() -> None:
    """PDF 解析器提取文本。"""
    import pymupdf

    from policymind.documents.parsers.pdf import PDFParser

    # 创建最小有效 PDF
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello PDF World", fontsize=12)
    pdf_bytes = doc.tobytes()
    doc.close()

    parser = PDFParser()
    result = parser.parse(pdf_bytes, filename="test.pdf")

    assert result.title is not None
    assert result.page_count == 1
    assert len(result.blocks) >= 1
    assert "Hello" in result.blocks[0].text


def test_docx_parser_extracts_paragraphs() -> None:
    """DOCX 解析器提取段落文本。"""
    from docx import Document

    from policymind.documents.parsers.docx import DocxParser

    buf = io.BytesIO()
    doc = Document()
    doc.add_heading("Policy Title", level=1)
    doc.add_paragraph("This is the first paragraph of the policy document.")
    doc.add_heading("Section A", level=2)
    doc.add_paragraph("Details about section A.")
    doc.save(buf)
    docx_bytes = buf.getvalue()

    parser = DocxParser()
    result = parser.parse(docx_bytes, filename="policy.docx")

    assert result.title == "Policy Title"
    assert len(result.blocks) >= 2
    texts = " ".join(b.text for b in result.blocks)
    assert "first paragraph" in texts


def test_xlsx_parser_extracts_table() -> None:
    """XLSX 解析器提取表格为 Markdown 格式。"""
    import openpyxl

    from policymind.documents.parsers.xlsx import XlsxParser

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Approvals"
    ws.append(["Name", "Amount", "Approver"])
    ws.append(["Alice", "5000", "Manager"])
    ws.append(["Bob", "12000", "Director"])

    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()
    wb.close()

    parser = XlsxParser()
    result = parser.parse(xlsx_bytes, filename="approvals.xlsx")

    assert result.title is not None
    assert len(result.blocks) >= 1
    assert "Alice" in result.blocks[0].text
    assert "Approver" in result.blocks[0].text


def test_parser_registry_all_formats() -> None:
    """ParserRegistry 注册并解析全部四种格式。"""
    from policymind.documents.parsers.docx import DocxParser
    from policymind.documents.parsers.markdown import MarkdownParser
    from policymind.documents.parsers.pdf import PDFParser
    from policymind.documents.parsers.registry import ParserRegistry
    from policymind.documents.parsers.xlsx import XlsxParser

    registry = ParserRegistry()
    registry.register(PDFParser())
    registry.register(DocxParser())
    registry.register(XlsxParser())
    registry.register(MarkdownParser())

    # 验证所有 MIME 都能解析
    mimes = [
        "application/pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "text/markdown",
    ]
    for mime in mimes:
        parser = registry.resolve(mime)
        assert parser is not None, f"No parser for {mime}"


def test_parser_registry_markdown_parse() -> None:
    """ParserRegistry 通过 MIME 解析 Markdown。"""
    from policymind.documents.parsers.markdown import MarkdownParser
    from policymind.documents.parsers.registry import ParserRegistry

    registry = ParserRegistry()
    registry.register(MarkdownParser())

    result = registry.parse(
        "text/markdown",
        content=b"# Hello\n\nWorld.",
        filename="test.md",
    )
    assert result.title == "Hello"


def test_parser_registry_unknown_mime_raises() -> None:
    """未知 MIME 类型抛出异常。"""
    import pytest

    from policymind.documents.parsers.registry import ParserRegistry

    registry = ParserRegistry()
    with pytest.raises(ValueError, match="parser"):
        registry.resolve("application/x-unknown")
